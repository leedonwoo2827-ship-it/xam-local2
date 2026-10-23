# -*- coding: utf-8 -*-
"""본문 무손실 검증 — 원본 HTML 과 생성된 PPTX 의 글자를 맞대 본다.

    python tools/verify_deck.py --html <_deck-export 또는 03 폴더> --pptx out

★ 왜 이게 필요한가. 이 파이프라인의 유일한 약속은 **원문을 그대로 옮긴다**는 것이다
  (LLM 을 안 쓰는 이유도 그것이다). 레이아웃을 손볼 때마다 문단 하나가 조용히
  사라질 수 있고 — 실제로 그랬다: 표 앞에 있던 `<p>1,000명을 대상으로…</p>` 가
  `_passage` 가 표만 뽑는 동안 없어졌고, 슬라이드를 눈으로 봐야 알았다.

★ 한 HTML 이 **pptx 여러 개**로 나온다(10문항/대단원마다 파일을 쪼갠다). 그래서
  파일 하나가 아니라 **그 원본에서 나온 파일 전부**의 글자를 이어서 본다.
  다른 원본의 파일까지 섞으면 한쪽에서 빠진 문장이 다른 쪽에 있어 통과해 버린다 —
  그래서 파일명 앞머리로 범위를 좁힌다.

★ 비교는 **정규화한 글자열의 포함 관계**다. 좌표·줄바꿈·블록 경계는 보지 않는다.
"""
from __future__ import annotations

import argparse
import glob
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from vendor.exam2deck import build, parse_html, parse_theory

_WS = re.compile(r"\s+")


def _norm(s: str) -> str:
    """공백을 지우고 눈에 안 보이는 차이를 없앤다.

    ★ 줄바꿈·들여쓰기는 레이아웃의 결과이므로 지운다. 다만 **글자는 하나도 지우지
      않는다** — 구두점을 빼고 비교하면 구두점이 사라진 사고를 못 잡는다.
    """
    return _WS.sub("", s or "")


def _pptx_text(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    buf = []
    for slide in prs.slides:
        for shp in slide.shapes:
            if shp.has_text_frame:
                buf.append(shp.text_frame.text)
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    for c in row.cells:
                        buf.append(c.text)
    return _norm("\n".join(buf))


def _pptx_pictures(path: str) -> int:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    return sum(1 for s in prs.slides for sh in s.shapes
               if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)


# ── 원본에서 "반드시 덱에 있어야 하는 조각" 을 뽑는다 ──────────────────────
def _blocks_fragments(blocks, tag: str):
    out = []
    for i, b in enumerate(blocks):
        if b.kind in ("passage", "passage_label", "pre", "expl", "bullets",
                      "heading", "q"):
            for j, p in enumerate(b.paras):
                out.append((f"{tag}{i}-{j}", p))
        elif b.kind == "table":
            for c in b.headers:
                out.append((f"{tag}표머리", c))
            for row in b.rows:
                for c in row:
                    out.append((f"{tag}표칸", c))
    return out


def _exam_expected(q):
    out = [("발문", " ".join(p for b in q.blocks if b.kind == "q"
                            for p in b.paras))]
    for cn, txt in q.choices():
        out.append((f"보기{cn}", txt))
    out.append(("정답번호", q.answer_no))
    if q.answer_text:
        out.append(("정답본문", q.answer_text))
    out += _blocks_fragments([b for b in q.blocks if b.kind != "q"], "지문")
    out += _blocks_fragments(q.expl, "해설")
    return [(k, v) for k, v in out if _norm(v)]


def _theory_expected(t):
    out = [("제목", f"{t.no} {t.title}".strip()), ("대단원", t.chapter)]
    out += _blocks_fragments(t.blocks, "본문")
    return [(k, v) for k, v in out if _norm(v)]


def _book_pptx(book: str, prefix: str) -> list:
    """`05/mNN-P/source/deck.pptx` 중 이 회차 것. prefix 는 "01회" 꼴이다."""
    import re as _re
    m = _re.match(r"(\d+)", prefix)
    if not m:
        return []
    rd = int(m.group(1))
    out = []
    root = os.path.join(book, "05")
    if not os.path.isdir(root):
        return []
    for name in sorted(os.listdir(root)):
        mm = _re.match(r"^m(\d{2})-(\d{1,2})$", name)
        if not mm or int(mm.group(1)) != rd:
            continue
        p = os.path.join(root, name, "source", "deck.pptx")
        if os.path.isfile(p):
            out.append((int(mm.group(2)), p))
    return [p for _n, p in sorted(out)]


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--html", default="", help="폴더 또는 .html 파일")
    ap.add_argument("--pptx", default="", help="산출 폴더(파일이 한곳에 모여 있을 때)")
    # ★ 앱 안에서는 pptx 가 한 폴더에 모여 있지 않다 — 번들마다
    #   `05/<번들>/source/deck.pptx` 다. --book 을 주면 거기서 찾는다.
    ap.add_argument("--book", default="", help="책 폴더. 주면 05/*/source/deck.pptx 를 본다")
    ap.add_argument("--show", type=int, default=8, help="불일치 표시 개수")
    a = ap.parse_args()

    if a.book and not a.html:
        a.html = os.path.join(a.book, "_deck-export")
    if not a.html:
        ap.error("--html 또는 --book 이 필요합니다")
    if not (a.pptx or a.book):
        ap.error("--pptx 또는 --book 이 필요합니다")

    srcs = (sorted(glob.glob(os.path.join(a.html, "*.html")))
            if os.path.isdir(a.html) else [a.html])
    total = miss = 0
    figs_want = figs_got = 0
    bad_srcs = []
    for src in srcs:
        theory = parse_theory.looks_like_theory(src)
        if theory:
            doc, _ = parse_theory.parse_file(src, "")
            prefix = build._safe(doc.title
                                 or os.path.splitext(os.path.basename(src))[0])
            units = doc.topics
            expect = _theory_expected
            want_fig = sum(1 for t in units for b in t.blocks
                           if b.kind == "figure")
            label = f"[이론] {doc.title}"
        else:
            rnd, _ = parse_html.parse_file(src, "")
            prefix = build._safe(rnd.round_no
                                 or os.path.splitext(os.path.basename(src))[0])
            units = rnd.questions
            expect = _exam_expected
            want_fig = sum(1 for q in units if q.figure)
            label = f"[문제] {rnd.round_no}"

        # 이 원본에서 나온 pptx 전부
        if a.book:
            files = _book_pptx(a.book, prefix)
        else:
            files = sorted(glob.glob(os.path.join(a.pptx, prefix + "*.pptx")))
        if not files:
            print(f"  {label}: PPTX 없음 (앞머리 '{prefix}')")
            bad_srcs.append(label)
            continue
        deck = "".join(_pptx_text(f) for f in files)
        got_fig = sum(_pptx_pictures(f) for f in files)

        bad = []
        n = 0
        for u in units:
            for tag, text in expect(u):
                n += 1
                total += 1
                if _norm(text) not in deck:
                    miss += 1
                    bad.append((getattr(u, "no", "?"), tag, text[:60]))
        figs_want += want_fig
        figs_got += got_fig
        ok = not bad and got_fig >= want_fig
        print(f"  {label}: 파일 {len(files)}개 · 조각 {n:5d}  누락 {len(bad):3d}  "
              f"도식 {got_fig}/{want_fig}  {'OK' if ok else '실패'}")
        for b in bad[:a.show]:
            print(f"      {b[0]} {b[1]}: {b[2]}")
        if not ok:
            bad_srcs.append(label)

    print(f"\n전체 조각 {total}  누락 {miss}  도식 {figs_got}/{figs_want}")
    if bad_srcs:
        print(f"실패: {bad_srcs}")
        return 1
    print("본문 무손실 · 도식 전량 확인")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
